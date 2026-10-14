"""Pure CSV -> transaction-row parsing. No FastAPI, no DB."""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.domain.classification import classify, is_transfer
from app.domain.enums import Ownership
from app.domain.errors import Invalid

MERCHANT_MAX = 200  # merchants.name is String(200)

FIELDS = {
    "date": ("date", "dato", "posted_at", "bokføringsdato", "transaksjonsdato"),
    "description": ("description", "beskrivelse", "tekst", "details", "transaction"),
    "amount": ("amount", "beløp", "belop", "sum"),
}
CURRENCY_NAMES = ("currency", "valuta")
KNOWN_HEADERS = {name for names in FIELDS.values() for name in names}


@dataclass(frozen=True)
class ParsedRow:
    row_number: int
    posted_at: date
    description: str
    amount: Decimal
    currency: str
    merchant: str
    category: str | None
    ownership: Ownership | None
    confidence: Decimal | None
    error: str | None
    is_transfer: bool = False


@dataclass(frozen=True)
class CsvRead:
    rows: list[ParsedRow]
    columns: list[str]  # label per column, for the mapping dropdowns
    mapping: dict[str, int | None]  # date / description / amount -> column index
    has_header: bool


def parse_date(value: str) -> date:
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Ugyldig dato: {value}")


def parse_amount(value: str) -> Decimal:
    value = value.strip().replace(" ", "").replace(" ", "")
    if not value:
        raise ValueError("Beløp mangler")
    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    elif "," in value:
        value = value.replace(",", ".")
    try:
        amount = Decimal(value)
    except InvalidOperation as exc:
        raise ValueError(f"Ugyldig beløp: {value}") from exc
    if not amount.is_finite():  # Decimal accepts "NaN" / "Infinity"
        raise ValueError(f"Ugyldig beløp: {value}")
    return amount


def _parses(fn, value: str) -> bool:
    try:
        fn(value)
        return True
    except ValueError:
        return False


def infer_columns(rows: list[list[str]]) -> dict[str, int]:
    """Guess date / description / amount columns from the values alone (partial result ok).

    ponytail: first all-numeric column wins amount (so date, text, amount, balance works);
    a file with the balance before the amount gets the wrong one. Upgrade path: let the
    user pick columns in the import preview.
    """
    sample = rows[:50]
    width = min(len(r) for r in sample)
    cols = [[r[i].strip() for r in sample if r[i].strip()] for i in range(width)]

    def parses(i, fn):
        return bool(cols[i]) and all(_parses(fn, v) for v in cols[i])

    date_i = next((i for i in range(width) if parses(i, parse_date)), None)
    amount_i = next((i for i in range(width) if i != date_i and parses(i, parse_amount)), None)
    text = [
        (sum(map(len, cols[i])) / len(cols[i]), i)
        for i in range(width)
        if i not in (date_i, amount_i) and cols[i] and not parses(i, parse_amount)
    ]
    found = {
        "date": date_i,
        "amount": amount_i,
        "description": max(text)[1] if text else None,
    }
    return {k: v for k, v in found.items() if v is not None}


def _delimiter(text: str) -> str:
    """First delimiter giving the same field count (> 1) on every line.

    Order matters: Norwegian exports use ';' with decimal commas, which fools csv.Sniffer.
    """
    head = [line for line in text.splitlines()[:20] if line.strip()]
    for d in ";\t,":
        counts = {len(next(csv.reader([line], delimiter=d))) for line in head}
        if len(counts) == 1 and counts != {1}:
            return d
    return ","


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")  # older Norwegian bank exports


def _parse_row(
    row_number: int,
    raw: list[str],
    cols: dict[str, int | None],
    currency_col: int | None,
) -> ParsedRow:
    def cell(index: int | None) -> str:
        return raw[index].strip() if index is not None and index < len(raw) else ""

    description = cell(cols["description"])
    currency = cell(currency_col) or "NOK"
    error = None
    try:
        posted_at = parse_date(cell(cols["date"]))
        amount = parse_amount(cell(cols["amount"]))
        if not description:
            raise ValueError("Beskrivelse mangler")
    except ValueError as exc:
        posted_at, amount, error = date.today(), Decimal(0), str(exc)

    classification = classify(description)
    return ParsedRow(
        row_number=row_number,
        posted_at=posted_at,
        description=description,
        amount=amount,
        currency=currency[:3].upper(),
        merchant=description[:MERCHANT_MAX],
        category=classification.category if classification else None,
        ownership=classification.ownership if classification else None,
        confidence=classification.confidence if classification else None,
        error=error,
        is_transfer=is_transfer(description),
    )


def read_csv(
    content: bytes,
    mapping: dict[str, int | None] | None = None,
    header: bool | None = None,
) -> CsvRead:
    """Parse a CSV. Columns come from the header names, else are inferred from the values;
    `mapping` (column indexes) and `header` override whatever was detected."""
    text = _decode(content)
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(text))
    lines = [r for r in reader if any(c.strip() for c in r)]
    if not lines:
        raise Invalid("CSV-filen er tom.")
    first, width = lines[0], max(map(len, lines))

    has_header = (
        header if header is not None else any(c.strip().casefold() in KNOWN_HEADERS for c in first)
    )
    body = lines[1:] if has_header else lines
    names = [c.strip().casefold() for c in first] if has_header else []

    def by_name(aliases):
        return next((names.index(a) for a in aliases if a in names), None)

    cols = {k: by_name(aliases) for k, aliases in FIELDS.items()}
    if None in cols.values() and body:  # infer only what the header didn't name
        guess = infer_columns(body)
        cols = {k: v if v is not None else guess.get(k) for k, v in cols.items()}
    for key, index in (mapping or {}).items():
        if index is not None:
            if not 0 <= index < width:
                raise Invalid("Ugyldig kolonne.")
            cols[key] = index

    def label(i: int) -> str:
        v = first[i].strip() if i < len(first) else ""
        return (v or f"Kolonne {i + 1}") if has_header else f"Kolonne {i + 1}: {v}"

    currency_col = by_name(CURRENCY_NAMES)
    start = 2 if has_header else 1
    rows = [_parse_row(n, raw, cols, currency_col) for n, raw in enumerate(body, start=start)]
    return CsvRead(rows, [label(i) for i in range(width)], cols, has_header)


def _xlsx_text(value) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return "" if value is None else str(value)


def xlsx_to_csv(content: bytes) -> bytes:
    """First worksheet of an XLSX workbook as UTF-8 CSV."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        workbook.close()
    except Exception as exc:
        raise Invalid(
            "Kunne ikke lese Excel-filen. Kontroller at det er en gyldig .xlsx-fil."
        ) from exc
    out = io.StringIO(newline="")
    csv.writer(out).writerows([_xlsx_text(v) for v in row] for row in rows)
    return out.getvalue().encode("utf-8")


def parse_csv(content: bytes) -> list[ParsedRow]:
    return read_csv(content).rows
